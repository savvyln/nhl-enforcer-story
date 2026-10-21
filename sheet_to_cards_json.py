"""
Turn the filled-in Player_Card_Images.xlsx into cards.json for the story.

Usage:  pip install openpyxl
        python sheet_to_cards_json.py Player_Card_Images.xlsx
Upload the resulting cards.json next to index.html. The story reads it on load,
and any card without an image falls back to a drawn placeholder card.
"""
import json, sys
from openpyxl import load_workbook

path = sys.argv[1] if len(sys.argv) > 1 else "Player_Card_Images.xlsx"
ws = load_workbook(path, data_only=False)["Cards"]
head = [c.value for c in ws[1]]
col = {name: i for i, name in enumerate(head)}

def cell(row, name):
    v = row[col[name]]
    return "" if v is None else str(v).strip()

out, linked = {}, 0
for row in ws.iter_rows(min_row=2, values_only=True):
    cid = cell(row, "Card ID")
    if not cid:
        continue
    num = cell(row, "Jersey Number")
    if num.endswith(".0"):
        num = num[:-2]
    rec = {
        "num": num,
        "img": cell(row, "Card Image URL (.png)"),
        "alt": cell(row, "Alt Text (optional)"),
        "pos": cell(row, "Position"),
        "pim": cell(row, "PIM* (5-season playoff total)"),
    }
    linked += bool(rec["img"])
    out[cid] = {k: v for k, v in rec.items() if v != ""}

json.dump(out, open("cards.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print(f"Wrote cards.json: {len(out)} cards, {linked} with images")
